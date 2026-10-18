import os, re, io, json, hmac, hashlib, secrets, smtplib, threading, time
from datetime import datetime, timedelta
from email.message import EmailMessage
from functools import wraps

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass
from flask import (Flask, request, session, redirect, render_template, flash,
                   jsonify, abort, send_file)
from flask_sqlalchemy import SQLAlchemy
from openpyxl import load_workbook, Workbook
from werkzeug.middleware.proxy_fix import ProxyFix
import matching

app = Flask(__name__)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1)
SECRET = os.environ.get("SECRET_KEY", "dev-change-me")
db_url = os.environ.get("DATABASE_URL") or "sqlite:///transfer.db"
if db_url.startswith("postgres://"):
    db_url = db_url.replace("postgres://", "postgresql://", 1)
app.config.update(SECRET_KEY=SECRET, SQLALCHEMY_DATABASE_URI=db_url,
                  SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax",
                  SESSION_COOKIE_SECURE=os.environ.get("HTTPS_ONLY") == "1",
                  MAX_CONTENT_LENGTH=2 * 1024 * 1024)
db = SQLAlchemy(app)
LOCK = threading.Lock()
NO_TRANSFER = "Currently no transfer needed"
GMAIL = re.compile(r"^[a-z0-9._+-]+@gmail\.com$")
GOOGLE_CLIENT_ID = os.environ.get("GOOGLE_CLIENT_ID", "")
OTP_ENABLED = os.environ.get("OTP_ENABLED") == "1"  # email OTP is off unless set to 1


def canon(email):
    """Gmail-only. Dots and +tags are ignored by Gmail, so strip them to stop duplicate accounts."""
    email = (email or "").strip().lower()
    if not GMAIL.match(email):
        return None
    local = email.split("@")[0].split("+")[0].replace(".", "")
    return local + "@gmail.com" if local else None


ADMINS = {canon(e) for e in os.environ.get("ADMIN_EMAILS", "").split(",") if canon(e)}
NOTIFY = os.environ.get("ADMIN_NOTIFY_EMAIL", "") or next(iter(ADMINS), "")

# ---------------------------------------------------------------- models
class Station(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(120), unique=True, nullable=False)

class Department(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(120), unique=True, nullable=False)

class Designation(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(120), unique=True, nullable=False)
    rank = db.Column(db.Integer, nullable=False)

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    email = db.Column(db.String(120), unique=True, nullable=False)
    name = db.Column(db.String(80), nullable=False)
    mobile = db.Column(db.String(15), nullable=False)
    station = db.Column(db.String(120), nullable=False)
    designation = db.Column(db.String(120), nullable=False)
    department = db.Column(db.String(120), nullable=False)
    choices = db.Column(db.Text, default="[]")
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

    @property
    def clist(self):
        try:
            v = json.loads(self.choices or "[]")
        except ValueError:  # "NONE" = user chose "Currently no transfer needed"
            return []
        return v if isinstance(v, list) else []

    @property
    def no_transfer(self):
        return self.choices == "NONE"

    @property
    def admin(self):
        return self.email in ADMINS

class Match(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    status = db.Column(db.String(10), default="pending")  # pending | confirmed
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    members = db.relationship("Member", backref="match", order_by="Member.pos",
                              cascade="all, delete-orphan")

class Member(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    match_id = db.Column(db.Integer, db.ForeignKey("match.id"), nullable=False)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), unique=True, nullable=False)
    pos = db.Column(db.Integer, nullable=False)
    response = db.Column(db.String(10))
    user = db.relationship("User")

class Block(db.Model):  # a rejected group is never proposed again
    id = db.Column(db.Integer, primary_key=True)
    key = db.Column(db.String(400), unique=True)

class Setting(db.Model):
    key = db.Column(db.String(20), primary_key=True)
    value = db.Column(db.String(20))

class Otp(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    email = db.Column(db.String(120), index=True)
    code_hash = db.Column(db.String(64))
    expires = db.Column(db.DateTime)
    tries = db.Column(db.Integer, default=0)
    created = db.Column(db.DateTime, default=datetime.utcnow)


class Recent(db.Model):  # anonymous "recent requirement" feed (one row per user)
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, unique=True)
    station = db.Column(db.String(120))
    designation = db.Column(db.String(120))
    wants = db.Column(db.String(600))
    at = db.Column(db.DateTime, default=datetime.utcnow)


def setting(key, default):
    s = db.session.get(Setting, key)
    return int(s.value) if s else default


_open = {"t": 0, "g": None}  # cached designation-free graph for "almost match" hints


def matching_on():
    return setting("ON", 1) == 1


def station_map():
    return {s.name.lower(): s.name for s in Station.query}


def ago(dt):
    s = int((datetime.utcnow() - dt).total_seconds())
    if s < 60:
        return "just now"
    if s < 3600:
        return f"{s // 60} min ago"
    if s < 86400:
        return f"{s // 3600} h ago"
    return f"{s // 86400} d ago"


def note_recent(u):
    r = Recent.query.filter_by(user_id=u.id).first() or Recent(user_id=u.id)
    r.station, r.designation, r.wants, r.at = u.station, u.designation, ", ".join(u.clist), datetime.utcnow()
    db.session.add(r)
    db.session.commit()


def near_matches(u):
    """Anonymous hints: chains that would exist if designation were ignored but are blocked by the K rule."""
    k, n = setting("K", 1), setting("N", 6)
    if _open["g"] is None or time.time() - _open["t"] > 120:
        ranks = {d.name: d.rank for d in Designation.query}
        taken = {m.user_id for m in Member.query}
        free = [x for x in User.query.all() if x.id not in taken and x.clist and x.designation in ranks]
        free.sort(key=lambda x: (not x.admin, x.created_at, x.id))
        people = [matching.P(x.id, x.station, x.clist, ranks[x.designation]) for x in free]
        adj, radj = matching.build_graph(people, 10**6)
        _open["g"] = (adj, radj, {p.id: p for p in people}, {x.id: x.designation for x in free})
        _open["t"] = time.time()
    adj, radj, pm, dname = _open["g"]
    if u.id not in adj:
        return []
    hints, blocked = [], set()
    for _ in range(5):
        cyc = matching.find_cycle(u.id, adj, radj, n, blocked)
        if not cyc:
            break
        blocked.add(frozenset(cyc))
        edges = [(cyc[i], cyc[(i + 1) % len(cyc)]) for i in range(len(cyc))]
        bad = [e for e in edges if abs(pm[e[0]].rank - pm[e[1]].rank) > k]
        if not bad:
            continue  # valid under K (e.g. a rejected group): not a designation problem
        prev = cyc[-1]  # the person who wants this user's station
        if (prev, u.id) in bad:
            h = (f"A {dname[prev]} person at {pm[prev].station} wants your station ({u.station}), "
                 f"but the designation difference is more than ±{k}.")
        else:
            h = f"A {len(cyc)}-person chain through your station exists, but one link has a designation difference above ±{k}."
        if h not in hints:
            hints.append(h)
    return hints[:3]

# ---------------------------------------------------------------- helpers
def send_mail(to, subject, body):
    host = os.environ.get("SMTP_HOST")
    if not host:
        print(f"[DEV MAIL] to={to} | {subject}\n{body}\n", flush=True)
        return True
    m = EmailMessage()
    m["From"] = os.environ.get("SMTP_FROM") or os.environ.get("SMTP_USER")
    m["To"], m["Subject"] = to, subject
    m.set_content(body)
    try:
        with smtplib.SMTP(host, int(os.environ.get("SMTP_PORT", 587)), timeout=15) as s:
            s.starttls()
            s.login(os.environ["SMTP_USER"], os.environ["SMTP_PASS"])
            s.send_message(m)
        return True
    except Exception as e:
        print("mail error:", e, flush=True)
        return False


def me():
    e = session.get("email")
    return User.query.filter_by(email=e).first() if e else None


def login_required(f):
    @wraps(f)
    def w(*a, **k):
        if not session.get("email"):
            return redirect("/login")
        return f(*a, **k)
    return w


def admin_required(f):
    @wraps(f)
    def w(*a, **k):
        if session.get("email") not in ADMINS:
            abort(403)
        return f(*a, **k)
    return w


@app.before_request
def csrf_protect():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    if request.method == "POST":
        tok = request.form.get("csrf") or request.headers.get("X-CSRF") or ""
        if not hmac.compare_digest(tok, session["csrf"]):
            abort(400)

app.jinja_env.globals.update(csrf=lambda: session.get("csrf", ""),
                             is_admin=lambda: session.get("email") in ADMINS)


def clean_mobile(raw):
    d = re.sub(r"\D", "", raw or "")
    if len(d) == 12 and d.startswith("91"):
        d = d[2:]
    return d if re.fullmatch(r"[6-9]\d{9}", d) else None


BASE = os.path.dirname(os.path.abspath(__file__))
EXCEL = os.environ.get("EXCEL_FILE") or os.path.join(BASE, "data.xlsx")
_xl = {"mtime": 0, "checked": 0, "error": "", "kept": 0}


def parse_excel(src):
    """Sheets: Stations (col A) | Designations (A name, B rank) | Departments (col A). Row 1 is a header."""
    wb = load_workbook(src, data_only=True)

    def rows(n):
        if n not in wb.sheetnames:
            raise ValueError(f"sheet '{n}' not found (found: {', '.join(wb.sheetnames)})")
        return [r for r in wb[n].iter_rows(min_row=2, values_only=True)
                if r and r[0] is not None and str(r[0]).strip()]

    stations = {str(r[0]).strip() for r in rows("Stations")}
    depts = {str(r[0]).strip() for r in rows("Departments")}
    desigs = {}
    for r in rows("Designations"):
        try:
            desigs[str(r[0]).strip()] = int(float(r[1]))
        except (TypeError, ValueError, IndexError):
            raise ValueError(f"Designations sheet: '{r[0]}' needs a number in column B (rank)")
    if not stations or not desigs or not depts:
        raise ValueError("Stations, Designations and Departments must each have at least one row")
    return stations, desigs, depts


def load_excel(src=EXCEL):
    """Make the dropdown tables match the Excel file. Items users still rely on are never deleted."""
    stations, desigs, depts = parse_excel(src)
    users = User.query.all()
    used_st = {u.station for u in users} | {c for u in users for c in u.clist}
    used_dg, used_dp = {u.designation for u in users}, {u.department for u in users}
    kept = 0
    for model, wanted, used in ((Station, stations, used_st), (Department, depts, used_dp)):
        have = set()
        for row in model.query.all():
            if row.name in wanted:
                have.add(row.name)
            elif row.name in used:
                kept += 1
            else:
                db.session.delete(row)
        for n in wanted - have:
            db.session.add(model(name=n))
    have = set()
    for d in Designation.query.all():
        if d.name in desigs:
            d.rank = desigs[d.name]
            have.add(d.name)
        elif d.name in used_dg:
            kept += 1
        else:
            db.session.delete(d)
    for n, rk in desigs.items():
        if n not in have:
            db.session.add(Designation(name=n, rank=rk))
    db.session.commit()
    _xl["kept"] = kept


def sync_excel(force=False):
    if not os.path.exists(EXCEL):
        _xl["error"] = f"{os.path.basename(EXCEL)} not found next to app.py"
        return
    m = os.path.getmtime(EXCEL)
    if not force and m == _xl["mtime"]:
        return
    _xl["mtime"] = m
    try:
        load_excel(EXCEL)
        _xl["error"] = ""
    except Exception as e:
        db.session.rollback()
        _xl["error"] = f"{type(e).__name__}: {e}"
        print("Excel load failed:", _xl["error"], flush=True)


@app.before_request
def excel_watch():  # picks up a replaced data.xlsx without a restart (checked at most every 30 s)
    if time.time() - _xl["checked"] > 30:
        _xl["checked"] = time.time()
        sync_excel()

# ---------------------------------------------------------------- matching
def run_matching(starts=None):
    """Find new matches. `starts`: user ids to search from (default: everyone free, by priority)."""
    if not matching_on():
        return 0
    k, n = setting("K", 1), setting("N", 6)
    ranks = {d.name: d.rank for d in Designation.query}
    taken = {m.user_id for m in Member.query}
    free = [u for u in User.query.all() if u.id not in taken and u.clist and u.designation in ranks]
    free.sort(key=lambda u: (not u.admin, u.created_at, u.id))
    people = [matching.P(u.id, u.station, u.clist, ranks[u.designation]) for u in free]
    adj, radj = matching.build_graph(people, k)
    blocked = {frozenset(int(x) for x in b.key.split(",")) for b in Block.query}
    order = [p.id for p in people]
    if starts is not None:
        s = set(starts)
        order = [i for i in order if i in s]
    used, found = set(), []
    for s in order:
        if s in used:
            continue
        cyc = matching.find_cycle(s, adj, radj, n, blocked, used)
        if cyc:
            m = Match()
            m.members = [Member(user_id=uid, pos=i) for i, uid in enumerate(cyc)]
            db.session.add(m)
            used.update(cyc)
            found.append(cyc)
    db.session.commit()
    _open["t"] = 0
    for cyc in found:
        for uid in cyc:
            u = db.session.get(User, uid)
            send_mail(u.email, "Mutual transfer: match found",
                      f"Hello {u.name},\n\nA mutual transfer match has been found for you. "
                      "Please log in to view the details and accept or reject it.")
    return len(found)


def drop_match(match, block=False):
    ids = sorted(m.user_id for m in match.members)
    if block and not Block.query.filter_by(key=",".join(map(str, ids))).first():
        db.session.add(Block(key=",".join(map(str, ids))))
    db.session.delete(match)
    db.session.commit()
    return ids

# ---------------------------------------------------------------- auth routes
@app.route("/login")
def login():
    return render_template("login.html", gid=GOOGLE_CLIENT_ID, otp=OTP_ENABLED)


@app.post("/logout")
def logout():
    session.clear()
    return redirect("/login")

_ip_hits = {}

@app.post("/otp/send")
def otp_send():
    if not OTP_ENABLED:
        abort(404)
    email = canon(request.form.get("email"))
    if not email:
        return jsonify(ok=False, msg="Only Gmail addresses (@gmail.com) are allowed."), 400
    now = time.time()
    hits = [t for t in _ip_hits.get(request.remote_addr, []) if now - t < 900]
    if len(hits) >= 10:
        return jsonify(ok=False, msg="Too many requests. Try again later."), 429
    _ip_hits[request.remote_addr] = hits + [now]
    recent = Otp.query.filter(Otp.email == email, Otp.created > datetime.utcnow() - timedelta(minutes=15)).all()
    if len(recent) >= 3 or any(o.created > datetime.utcnow() - timedelta(seconds=60) for o in recent):
        return jsonify(ok=False, msg="Please wait a minute before requesting another OTP."), 429
    code = f"{secrets.randbelow(10**6):06d}"
    db.session.add(Otp(email=email, code_hash=hashlib.sha256((SECRET + email + code).encode()).hexdigest(),
                       expires=datetime.utcnow() + timedelta(minutes=10)))
    db.session.commit()
    send_mail(email, "Your login OTP", f"Your OTP is {code}. It is valid for 10 minutes.")
    return jsonify(ok=True, msg="OTP sent. Check your Gmail inbox.")


@app.post("/otp/verify")
def otp_verify():
    if not OTP_ENABLED:
        abort(404)
    email, code = canon(request.form.get("email")), (request.form.get("otp") or "").strip()
    o = email and Otp.query.filter_by(email=email).order_by(Otp.id.desc()).first()
    if not o or o.expires < datetime.utcnow() or o.tries >= 5:
        return jsonify(ok=False, msg="OTP expired. Request a new one."), 400
    good = hashlib.sha256((SECRET + email + code).encode()).hexdigest()
    if not hmac.compare_digest(good, o.code_hash):
        o.tries += 1
        db.session.commit()
        return jsonify(ok=False, msg="Incorrect OTP."), 400
    Otp.query.filter_by(email=email).delete()
    db.session.commit()
    session.clear()
    session["email"] = email
    return jsonify(ok=True)


@app.post("/google")
def google_login():
    from google.oauth2 import id_token
    from google.auth.transport import requests as grequests
    try:
        info = id_token.verify_oauth2_token((request.get_json(silent=True) or {}).get("credential", ""),
                                            grequests.Request(), GOOGLE_CLIENT_ID)
    except Exception:
        return jsonify(ok=False, msg="Google sign-in failed."), 400
    email = canon(info.get("email"))
    if not email or not info.get("email_verified"):
        return jsonify(ok=False, msg="Only Gmail addresses are allowed."), 400
    session.clear()
    session["email"] = email
    return jsonify(ok=True)

# ---------------------------------------------------------------- user routes
@app.route("/")
def index():
    if not session.get("email"):
        return redirect("/login")
    return redirect("/dashboard" if me() else "/register")


@app.route("/how-it-works")
def how():
    return render_template("how.html", K=setting("K", 1), N=setting("N", 6))


@app.route("/register", methods=["GET", "POST"])
@login_required
def register():
    if me():
        return redirect("/dashboard")
    if request.method == "POST":
        f = request.form
        name, mob = f.get("full_name", "").strip(), clean_mobile(f.get("mobile"))
        st = station_map().get(f.get("station", "").strip().lower())
        ok = (2 <= len(name) <= 80 and mob and st
              and Designation.query.filter_by(name=f.get("designation")).first()
              and Department.query.filter_by(name=f.get("department")).first())
        if not ok:
            flash("Please fill all fields correctly (10-digit Indian mobile number).", "error")
        else:
            db.session.add(User(email=session["email"], name=name, mobile=mob, station=st,
                                designation=f["designation"], department=f["department"]))
            db.session.commit()
            flash("Registered. Now add your preferred stations.", "ok")
            return redirect("/dashboard")
    return render_template("register.html", stations=Station.query.order_by(Station.name).all(),
                           designations=Designation.query.order_by(Designation.rank).all(),
                           departments=Department.query.order_by(Department.name).all())


@app.route("/dashboard")
@login_required
def dashboard():
    u = me()
    if not u:
        return redirect("/register")
    mem = Member.query.filter_by(user_id=u.id).first()
    chain = status = my_resp = None
    if mem:
        ms = mem.match.members
        chain = [dict(name=m.user.name, designation=m.user.designation, department=m.user.department,
                      mobile=m.user.mobile, email=m.user.email, frm=m.user.station,
                      to=ms[(i + 1) % len(ms)].user.station, resp=m.response, me=m.user_id == u.id)
                 for i, m in enumerate(ms)]
        status, my_resp = mem.match.status, mem.response
    recent = [dict(station=r.station, designation=r.designation, wants=r.wants, ago=ago(r.at))
              for r in Recent.query.filter(Recent.user_id != u.id).order_by(Recent.at.desc()).limit(4)]
    on = matching_on()
    hints = near_matches(u) if (not mem and u.clist and on) else []
    return render_template("dashboard.html", u=u, chain=chain, status=status, my_resp=my_resp,
                           stations=Station.query.order_by(Station.name).all(), choices=u.clist,
                           designations=Designation.query.order_by(Designation.rank).all(),
                           departments=Department.query.order_by(Department.name).all(),
                           recent=recent, hints=hints, on=on, no_transfer_text=NO_TRANSFER)


@app.post("/mobile")
@login_required
def set_mobile():
    u, mob = me(), clean_mobile(request.form.get("mobile"))
    if u and mob:
        u.mobile = mob
        db.session.commit()
        flash("Mobile number updated.", "ok")
    else:
        flash("Enter a valid 10-digit mobile number.", "error")
    return redirect("/dashboard")


@app.post("/profile")
@login_required
def update_profile():
    u = me()
    if not u:
        return redirect("/register")
    f = request.form
    st = station_map().get(f.get("station", "").strip().lower())
    name, mob = f.get("full_name", "").strip(), clean_mobile(f.get("mobile"))
    if not (2 <= len(name) <= 80 and mob and st and Designation.query.filter_by(name=f.get("designation")).first()
            and Department.query.filter_by(name=f.get("department")).first()):
        flash("Please fill all fields correctly.", "error")
        return redirect("/dashboard")
    moved = (st, f["designation"]) != (u.station, u.designation)
    with LOCK:
        ids = [u.id]
        mem = Member.query.filter_by(user_id=u.id).first()
        if mem and moved:
            ids = drop_match(mem.match)  # old match no longer valid
        u.name, u.mobile, u.station, u.designation, u.department = name, mob, st, f["designation"], f["department"]
        if not u.no_transfer:
            u.choices = json.dumps([c for c in u.clist if c != st])
        r = Recent.query.filter_by(user_id=u.id).first()
        if r:
            r.station, r.designation = st, u.designation
        db.session.commit()
        _open["t"] = 0
        if moved:
            run_matching(ids)
    flash("Profile updated.", "ok")
    return redirect("/dashboard")


@app.post("/choices")
@login_required
def set_choices():
    u = me()
    if not u:
        return redirect("/register")
    smap = station_map()
    raw = [v.strip() for v in request.form.getlist("c") if v.strip()]
    vals = [smap.get(v.lower()) for v in raw]
    if Member.query.filter_by(user_id=u.id).first():
        flash("Respond to your current match before changing choices.", "error")
    elif raw and raw[0].lower() == NO_TRANSFER.lower():
        u.choices = "NONE"
        Recent.query.filter_by(user_id=u.id).delete()
        db.session.commit()
        _open["t"] = 0
        flash("Saved. You are marked as not needing a transfer right now. You can change this anytime.", "ok")
    elif not vals or len(set(vals)) != len(vals) or len(vals) > 5 or None in vals or u.station in vals:
        flash("Pick 1 to 5 different stations from the list (first is mandatory, not your own station).", "error")
    else:
        u.choices = json.dumps(vals)
        db.session.commit()
        note_recent(u)
        with LOCK:
            run_matching([u.id])
        flash("Choices saved.", "ok")
    return redirect("/dashboard")


@app.post("/match/respond")
@login_required
def respond():
    u, action = me(), request.form.get("action")
    with LOCK:
        mem = Member.query.filter_by(user_id=u.id).first() if u else None
        if not mem or mem.match.status != "pending":
            flash("No pending match.", "error")
        elif action == "accept":
            mem.response = "accept"
            _open["t"] = 0
            if all(x.response == "accept" for x in mem.match.members):
                mem.match.status = "confirmed"
                for x in mem.match.members:
                    send_mail(x.user.email, "Mutual transfer confirmed",
                              "All members accepted. Log in to see contact details.")
            db.session.commit()
        elif action == "reject":
            ids = drop_match(mem.match, block=True)
            run_matching(ids)
            flash("Match rejected. You stay in the pool for other matches.", "ok")
    return redirect("/dashboard")

# ---------------------------------------------------------------- admin
@app.route("/admin")
@admin_required
def admin():
    q = request.args.get("q", "").strip()
    users = User.query
    if q:
        users = users.filter(db.or_(User.name.ilike(f"%{q}%"), User.email.ilike(f"%{q}%")))
    users = users.order_by(User.created_at).limit(100).all()
    edit = db.session.get(User, request.args.get("edit", type=int) or 0)
    stats = dict(users=User.query.count(), pending=Match.query.filter_by(status="pending").count(),
                 confirmed=Match.query.filter_by(status="confirmed").count(),
                 with_choices=User.query.filter(~User.choices.in_(["[]", "NONE"])).count())
    return render_template("admin.html", users=users, edit=edit, stats=stats, q=q,
                           K=setting("K", 1), N=setting("N", 6), on=matching_on(), xl_error=_xl["error"], xl_kept=_xl["kept"], matches=Match.query.all(),
                           stations=Station.query.order_by(Station.name).all(),
                           designations=Designation.query.order_by(Designation.rank).all(),
                           departments=Department.query.order_by(Department.name).all())


@app.post("/admin/settings")
@admin_required
def admin_settings():
    k, n = request.form.get("K", type=int), request.form.get("N", type=int)
    if k is None or n is None or not (0 <= k <= 20 and 2 <= n <= 30):
        flash("K must be 0-20 and N 2-30.", "error")
    else:
        for key, val in (("K", k), ("N", n)):
            s = db.session.get(Setting, key) or Setting(key=key)
            s.value = str(val)
            db.session.add(s)
        db.session.commit()
        with LOCK:
            found = run_matching()
        flash(f"Saved. {found} new match(es) found.", "ok")
    return redirect("/admin")


@app.post("/admin/toggle")
@admin_required
def admin_toggle():
    on = 0 if matching_on() else 1
    s = db.session.get(Setting, "ON") or Setting(key="ON")
    s.value = str(on)
    db.session.add(s)
    db.session.commit()
    if on:
        with LOCK:
            flash(f"Matching is ON. {run_matching()} new match(es) found.", "ok")
    else:
        flash("Matching is OFF. Users can still save choices; matching resumes when you switch it on.", "ok")
    return redirect("/admin")


@app.post("/admin/reload")
@admin_required
def admin_reload():
    sync_excel(force=True)
    flash("Reloaded from data.xlsx." if not _xl["error"] else f"Excel problem: {_xl['error']}",
          "ok" if not _xl["error"] else "error")
    return redirect("/admin")


@app.post("/admin/rematch")
@admin_required
def admin_rematch():
    with LOCK:
        flash(f"{run_matching()} new match(es) found.", "ok")
    return redirect("/admin")


@app.post("/admin/clearblocks")
@admin_required
def admin_clearblocks():
    Block.query.delete()
    db.session.commit()
    flash("Rejected-group blocks cleared.", "ok")
    return redirect("/admin")


@app.post("/admin/user/<int:uid>")
@admin_required
def admin_user(uid):
    u = db.session.get(User, uid) or abort(404)
    f = request.form
    new = dict(name=f.get("full_name", "").strip(), mobile=clean_mobile(f.get("mobile")),
               station=f.get("station"), designation=f.get("designation"), department=f.get("department"))
    if not (2 <= len(new["name"]) <= 80 and new["mobile"] and Station.query.filter_by(name=new["station"]).first()
            and Designation.query.filter_by(name=new["designation"]).first()
            and Department.query.filter_by(name=new["department"]).first()):
        flash("Invalid values.", "error")
        return redirect(f"/admin?edit={uid}")
    changes = [f"{k}: {getattr(u, k)} -> {v}" for k, v in new.items() if getattr(u, k) != v]
    if changes:
        with LOCK:
            ids = [u.id]
            mem = Member.query.filter_by(user_id=u.id).first()
            if mem and any(not c.startswith(("name", "mobile")) for c in changes):
                ids = drop_match(mem.match)  # position changed: old match no longer valid
            for k, v in new.items():
                setattr(u, k, v)
            if u.station in u.clist:
                u.choices = json.dumps([c for c in u.clist if c != u.station])
            db.session.commit()
            run_matching(ids)
        send_mail(NOTIFY, f"Admin edit: {u.email}", f"Changes made by {session['email']}:\n" + "\n".join(changes))
    flash("User updated." if changes else "No changes.", "ok")
    return redirect("/admin")


@app.route("/admin/export")
@admin_required
def admin_export():
    wb = Workbook()
    ws = wb.active
    ws.title = "Matches"
    ws.append(["Match", "Status", "Name", "Email", "Mobile", "Designation", "Department", "From", "To", "Response"])
    for m in Match.query.all():
        for i, x in enumerate(m.members):
            ws.append([m.id, m.status, x.user.name, x.user.email, x.user.mobile, x.user.designation,
                       x.user.department, x.user.station, m.members[(i + 1) % len(m.members)].user.station,
                       x.response or "pending"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="matches.xlsx")


with app.app_context():
    db.create_all()
    sync_excel(force=True)

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=os.environ.get("DEBUG") == "1")
